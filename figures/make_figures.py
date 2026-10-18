#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Sinh toàn bộ hình vẽ của báo cáo KTCN từ SỐ LIỆU THỰC TẾ.

Nguồn dữ liệu:
  - 'Cấp Giấy Giới Thiệu Thực Tập (Responses).xlsx' : 361 hồ sơ đã được hệ thống xử lý
    (cột TRẠNG THÁI GỬI GGT, LINK PDF GGT, THỜI GIAN XỬ LÝ GGT, Kết quả đối soát MST).
  - ggt-cntt-auto/sources/Code.gs                     : cấu hình CONFIG mới nhất.
  - ggt-cntt-auto/sources/test_suite.js               : 93 ca kiểm thử (node test_suite.js).

Cách chạy:  python figures/make_figures.py
"""
import collections
import datetime
import os
import re
import statistics
import sys

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch, Rectangle
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIGDIR = os.path.join(ROOT, "figures")
XLSX = os.path.join(ROOT, "Cấp Giấy Giới Thiệu Thực Tập (Responses).xlsx")

# ---------------------------------------------------------------- style
NAVY = "#1F4E79"
BLUE = "#2E75B6"
GREEN = "#2E7D32"
GREEN_L = "#E3F1E4"
RED = "#C62828"
RED_L = "#FBE4E4"
ORANGE = "#E65100"
GREY = "#6B7280"
YELLOW_L = "#FFF6DA"
PURPLE_L = "#EFE6F7"
TEAL_L = "#E2F1F1"

plt.rcParams.update({
    "font.family": "DejaVu Sans",
    "font.size": 12,
    "axes.edgecolor": "#8A8A8A",
    "axes.labelcolor": "#222222",
    "text.color": "#222222",
    "figure.dpi": 200,
})


def titles(ax, title, subtitle=None):
    ax.set_title(title, fontsize=16, fontweight="bold", color=NAVY, pad=34 if subtitle else 10)
    if subtitle:
        ax.text(0.5, 1.015, subtitle, ha="center", va="bottom", transform=ax.transAxes,
                fontsize=11.5, color=GREY)


def save(fig, name):
    path = os.path.join(FIGDIR, name)
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    print("  ->", os.path.relpath(path, ROOT))


def box(ax, x, y, w, h, text, fc, ec, fs=11.5, weight="bold", tc="#1A1A1A", radius=0.03):
    ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle=f"round,pad=0.004,rounding_size={radius}",
                                linewidth=1.8, edgecolor=ec, facecolor=fc))
    ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=fs,
            fontweight=weight, color=tc, linespacing=1.5)


def arrow(ax, p1, p2, color=NAVY, style="-|>", lw=2.2, rad=0.0, ls="-"):
    ax.add_patch(FancyArrowPatch(p1, p2, arrowstyle=style, mutation_scale=18,
                                 linewidth=lw, color=color, linestyle=ls,
                                 connectionstyle=f"arc3,rad={rad}"))


# ---------------------------------------------------------------- data
class Data:
    def __init__(self, path):
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
        headers = [c.value for c in ws[1]]
        self.idx = {h: i for i, h in enumerate(headers)}
        self.rows = [r for r in ws.iter_rows(min_row=2, values_only=True)
                     if isinstance(r[0], datetime.datetime)]
        S = self.idx["TRẠNG THÁI GỬI GGT"]
        T = self.idx["THỜI GIAN XỬ LÝ GGT"]
        self.success = [r for r in self.rows if str(r[S]).strip() == "Đã tạo PDF"]
        self.errors = [r for r in self.rows if str(r[S]).startswith("Lỗi")]
        times = sorted(r[T] for r in self.rows if isinstance(r[T], datetime.datetime))
        self.times = times
        # các "đợt chạy" liên tục: khoảng cách giữa 2 hồ sơ liên tiếp <= 120 giây
        runs, cur = [], [times[0]]
        for a, b in zip(times, times[1:]):
            if (b - a).total_seconds() > 120:
                runs.append(cur); cur = [b]
            else:
                cur.append(b)
        runs.append(cur)
        self.runs = runs
        self.active_seconds = sum((r[-1] - r[0]).total_seconds() for r in runs)
        deltas = [d for d in ((b - a).total_seconds() for a, b in zip(times, times[1:])) if 0 < d < 60]
        self.median_gap = statistics.median(deltas)

    @property
    def total(self):
        return len(self.rows)

    @property
    def n_ok(self):
        return len(self.success)

    @property
    def n_err(self):
        return len(self.errors)

    def status_text(self, r):
        return str(r[self.idx["TRẠNG THÁI GỬI GGT"]])

    def note(self, r):
        return str(r[self.idx["Kết quả đối soát MST"]])

    def error_class(self, r):
        s = self.status_text(r)
        if "ngừng hoạt động" in s:
            return "MST thuộc doanh nghiệp\nđã ngừng hoạt động"
        if "không tồn tại trên hệ thống Thuế" in s:
            return "MST không tồn tại\ntrên hệ thống Thuế"
        if "tên công ty không hợp lệ" in s:
            return "MST đúng nhưng tên\ndoanh nghiệp không khớp"
        if "thiếu Tên công ty" in s:
            return "Thiếu tên doanh nghiệp"
        if "Số điện thoại không hợp lệ" in s:
            return "Số điện thoại sai định dạng"
        return "MST sai cú pháp / để trống"

    def match_class(self, r):
        s = self.note(r)
        if "Khớp chính xác" in s:
            return "Khớp chính xác tuyệt đối"
        if "tương đồng cao" in s:
            return "Tương đồng cao (Weighted Dice ≥ 50%)"
        if "viết tắt" in s:
            return "Khớp tên viết tắt thương hiệu"
        if "công ty tổng" in s:
            return "Khớp công ty tổng của chi nhánh"
        return "Khớp tên tiếng nước ngoài"

    def courses(self):
        c = collections.Counter()
        for r in self.success:
            v = str(r[self.idx["Tên Học Phần"]] or "").strip()
            if v.startswith("Kiến Tập Công Nghiệp, Tập Sự"):
                v = "KTCN + TSNN"
            elif v.startswith("Kiến Tập Công Nghiệp, Dự"):
                v = "KTCN + Dự án CNTT 1"
            elif v.startswith("Kiến Tập Công Nghiệp"):
                v = "Kiến tập công nghiệp"
            elif v.startswith("Tập Sự"):
                v = "Tập sự nghề nghiệp"
            elif v.startswith("Dự Án"):
                v = "Dự án CNTT 1"
            else:
                v = "Chưa đăng ký"
            c[v] += 1
        return c

    def companies(self, top=10):
        c = collections.Counter(str(r[self.idx["Tên Doanh Nghiệp"]]).strip().upper()
                                for r in self.success)
        return c.most_common(top), len(c)


# ---------------------------------------------------------------- figures
def fig_ket_qua(d):
    """Hình: tỷ lệ xử lý + phân loại 31 lỗi."""
    err = collections.Counter(d.error_class(r) for r in d.errors)
    fig = plt.figure(figsize=(16, 7.2))
    gs = fig.add_gridspec(1, 2, width_ratios=[1, 1.25], wspace=0.28)

    ax = fig.add_subplot(gs[0, 0])
    vals = [d.n_ok, d.n_err]
    labels = [f"Đã tạo PDF thành công\n{d.n_ok} hồ sơ ({d.n_ok/d.total*100:.1f}%)",
              f"Bị chặn và ghi lỗi\n{d.n_err} hồ sơ ({d.n_err/d.total*100:.1f}%)"]
    wedges, _ = ax.pie(vals, colors=[GREEN, RED], startangle=90,
                       wedgeprops=dict(width=0.42, edgecolor="white", linewidth=2))
    ax.text(0, 0.08, f"{d.total}", ha="center", va="center", fontsize=30,
            fontweight="bold", color=NAVY)
    ax.text(0, -0.22, "hồ sơ thực tế\n(trong Sheet của Khoa)", ha="center", va="center",
            fontsize=12, color=GREY)
    ax.annotate(labels[0], xy=(-0.72, -0.72), xytext=(-1.62, -1.28), fontsize=12.5,
                fontweight="bold", color=GREEN,
                arrowprops=dict(arrowstyle="-", color=GREEN, lw=1.4))
    ax.annotate(labels[1], xy=(0.42, 0.72), xytext=(0.30, 1.28), fontsize=12.5,
                fontweight="bold", color=RED,
                arrowprops=dict(arrowstyle="-", color=RED, lw=1.4))
    ax.set_title("KẾT QUẢ XỬ LÝ THỰC TẾ", fontsize=16, fontweight="bold", color=NAVY, pad=22)
    ax.set_xlim(-1.9, 1.5); ax.set_ylim(-1.5, 1.6)

    ax2 = fig.add_subplot(gs[0, 1])
    items = err.most_common()
    ys = range(len(items))
    ax2.barh(list(ys), [v for _, v in items], color=RED, height=0.55, edgecolor="#8E1F1F")
    for y, (k, v) in zip(ys, items):
        ax2.text(v + 0.18, y, f"{v} ca  ({v/d.n_err*100:.1f}%)", va="center",
                 fontsize=12, fontweight="bold", color=RED)
    ax2.set_yticks(list(ys))
    ax2.set_yticklabels([k for k, _ in items], fontsize=11.5)
    ax2.invert_yaxis()
    ax2.set_xlim(0, max(v for _, v in items) + 2.6)
    ax2.set_xlabel("Số hồ sơ bị hệ thống chặn và ghi lỗi", fontsize=12)
    ax2.grid(axis="x", alpha=0.25, linestyle=":")
    for s in ("top", "right"):
        ax2.spines[s].set_visible(False)
    titles(ax2, "PHÂN LOẠI 31 HỒ SƠ BỊ CHẶN",
           "Nguồn: cột “TRẠNG THÁI GỬI GGT” của Sheet, đợt chạy 17–18/10/2026")
    save(fig, "hinh_5_6_bieu_do_ket_qua.png")


def fig_thoi_gian(d, manual_min_per_record=5.0):
    """Hình: so sánh thời gian thủ công và tự động (số liệu đo thật)."""
    auto_min = d.active_seconds / 60.0
    manual_min = d.total * manual_min_per_record
    per_auto = d.active_seconds / d.total
    per_manual = manual_min_per_record * 60
    save_pct = (1 - auto_min / manual_min) * 100

    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(16, 7.0))

    bars = ax1.bar(["Xử lý thủ công\n(ước lượng từ khảo sát)",
                    "GGT-CNTT Auto\n(đo từ thời gian chạy thật)"],
                   [manual_min, auto_min], color=["#F4801F", "#1E8E5A"], width=0.52,
                   edgecolor="#4D4D4D")
    for b, v in zip(bars, [manual_min, auto_min]):
        ax1.text(b.get_x() + b.get_width() / 2, v + 45, f"{v:,.0f} phút", ha="center",
                 fontsize=15, fontweight="bold", color="#1A1A1A")
    ax1.annotate("", xy=(1, auto_min + 8), xytext=(0, manual_min),
                 arrowprops=dict(arrowstyle="-|>", color=GREEN, lw=2.2,
                                 connectionstyle="arc3,rad=-0.25"))
    ax1.text(0.52, manual_min * 0.55, f"Tiết kiệm\n{save_pct:.1f}% thời gian",
             ha="center", fontsize=14, fontweight="bold", color=GREEN)
    ax1.set_ylabel(f"Tổng thời gian xử lý {d.total} hồ sơ (phút)", fontsize=12)
    ax1.set_ylim(0, manual_min * 1.15)
    ax1.grid(axis="y", alpha=0.25, linestyle=":")
    titles(ax1, f"TỔNG THỜI GIAN XỬ LÝ {d.total} HỒ SƠ")
    ax1.text(0.5, -0.17, f"Thủ công: {manual_min_per_record:.0f} phút/hồ sơ (khảo sát Văn phòng Khoa). "
                          f"Tự động: {auto_min:.1f} phút cho cả {d.total} hồ sơ, đo từ 4 đợt chạy liên tục.",
             transform=ax1.transAxes, ha="center", fontsize=10.5, color=GREY)

    labels2 = ["Thời gian/1 hồ sơ", "Hồ sơ xử lý được\ntrong 1 phút", "Tệp PDF tạo được\ntrong 1 phút"]
    manual_vals = [per_manual, 60 / per_manual, 60 / per_manual]
    auto_vals = [per_auto, d.total / auto_min, d.n_ok / auto_min]
    x = range(3)
    w = 0.36
    b1 = ax2.bar([i - w / 2 for i in x], manual_vals, w, label="Thủ công (ước lượng)",
                 color="#F4801F", edgecolor="#4D4D4D")
    b2 = ax2.bar([i + w / 2 for i in x], auto_vals, w, label="GGT-CNTT Auto (đo thật)",
                 color="#1E8E5A", edgecolor="#4D4D4D")
    for bars in (b1, b2):
        for b in bars:
            v = b.get_height()
            ax2.text(b.get_x() + b.get_width() / 2, v * 1.08,
                     f"{v:.1f}" if v < 100 else f"{v:,.0f}",
                     ha="center", fontsize=12, fontweight="bold", color="#1A1A1A")
    ax2.set_yscale("log")
    ax2.set_ylim(0.05, 2000)
    ax2.set_xticks(list(x)); ax2.set_xticklabels(labels2, fontsize=11.5)
    ax2.set_ylabel("Giá trị (thang logarit)", fontsize=12)
    ax2.legend(fontsize=11.5, loc="upper right", framealpha=0.95)
    ax2.grid(axis="y", alpha=0.25, linestyle=":")
    titles(ax2, "CHỈ SỐ HIỆU NĂNG VẬN HÀNH",
           f"Hồ sơ trung bình: {per_manual:.0f} giây → {per_auto:.1f} giây (giảm {save_pct:.1f}%)")
    save(fig, "hinh_5_9_bieu_do_thoi_gian.png")
    return dict(manual_min=manual_min, auto_min=auto_min, per_auto=per_auto,
                save_pct=save_pct, throughput=d.total / auto_min)


def fig_dashboard(d):
    """Hình: Dashboard tổng hợp với số liệu thực tế."""
    fig = plt.figure(figsize=(15, 8.2))
    fig.suptitle("DASHBOARD GGT — BẢNG TỔNG HỢP SINH TỪ SỐ LIỆU THỰC TẾ",
                 fontsize=17, fontweight="bold", color=NAVY, y=0.98)
    ax = fig.add_axes([0, 0, 1, 1]); ax.axis("off")

    cards = [("Tổng số hồ sơ", d.total, NAVY),
             ("Đã tạo PDF", d.n_ok, GREEN),
             ("Lỗi (đã ghi rõ nguyên nhân)", d.n_err, RED),
             ("Đang xử lý", 0, ORANGE),
             ("Chờ duyệt", 0, BLUE)]
    w, gap = 0.172, 0.024
    for i, (label, value, color) in enumerate(cards):
        x = 0.03 + i * (w + gap)
        box(ax, x, 0.70, w, 0.16, "", "#FFFFFF", color)
        ax.text(x + w / 2, 0.815, f"{value}", ha="center", fontsize=30,
                fontweight="bold", color=color)
        ax.text(x + w / 2, 0.735, label, ha="center", fontsize=10.5, color="#333333")

    ax2 = fig.add_axes([0.06, 0.12, 0.44, 0.46])
    n = [d.n_ok, d.n_err, 0, 0]
    ax2.bar(["Đã tạo PDF", "Lỗi", "Đang xử lý", "Chờ duyệt"], n,
            color=[GREEN, RED, ORANGE, BLUE], width=0.55, edgecolor="#4D4D4D")
    for i, v in enumerate(n):
        ax2.text(i, v + 6, str(v), ha="center", fontsize=14, fontweight="bold")
    ax2.set_ylim(0, max(n) * 1.2)
    ax2.set_ylabel("Số hồ sơ", fontsize=11.5)
    ax2.grid(axis="y", alpha=0.25, linestyle=":")
    for s in ("top", "right"):
        ax2.spines[s].set_visible(False)
    titles(ax2, "TRẠNG THÁI HỒ SƠ",
           f"Cập nhật lúc {d.times[-1].strftime('%d/%m/%Y %H:%M:%S')}")

    ax3 = fig.add_axes([0.56, 0.12, 0.41, 0.46]); ax3.axis("off")
    ax3.set_title("CÔNG THỨC TÍNH TRÊN SHEET “Dashboard GGT”", fontsize=13.5,
                  fontweight="bold", color=NAVY, pad=12)
    rows = [("Tổng số hồ sơ", "=COUNTA('Form Responses 1'!B2:B)"),
            ("Đã tạo PDF", '=COUNTIF(...;"Đã tạo PDF")+COUNTIF(...;"Đã gửi")'),
            ("Đang xử lý", '=COUNTIF(...;"Processing")'),
            ("Chờ duyệt", '=COUNTIF(...;"Chờ duyệt")'),
            ("Lỗi", '=COUNTIF(...;"Lỗi:*")')]
    for i, (a, b) in enumerate(rows):
        y = 0.82 - i * 0.155
        ax3.add_patch(Rectangle((0.0, y - 0.10), 1.0, 0.145, facecolor="#F4F6F8",
                                edgecolor="#C9D2DA"))
        ax3.text(0.03, y - 0.028, a, fontsize=11.5, fontweight="bold", va="center")
        ax3.text(0.34, y - 0.028, b, fontsize=10.5, family="DejaVu Sans Mono",
                 color="#0B4F6C", va="center")
    ax3.text(0.0, -0.02, "Hàm refreshDashboard() cập nhật lại bảng này bằng một lần bấm menu,\n"
                         "không cần mở công thức tính tay như trước.",
             fontsize=11, color=GREY)
    save(fig, "hinh_5_10_dashboard.png")


def fig_doi_soat(d):
    """Hình: các phương thức đối soát đã xác nhận hồ sơ."""
    c = collections.Counter(d.match_class(r) for r in d.success)
    items = c.most_common()
    fig, ax = plt.subplots(figsize=(15, 6.6))
    ys = list(range(len(items)))
    colors = [GREEN, "#43A047", ORANGE, BLUE, "#7E57C2"]
    ax.barh(ys, [v for _, v in items], color=colors[:len(items)], height=0.55,
            edgecolor="#4D4D4D")
    for y, (k, v) in zip(ys, items):
        ax.text(v + d.n_ok * 0.006, y, f"{v} hồ sơ  ({v/d.n_ok*100:.1f}%)", va="center",
                fontsize=12.5, fontweight="bold", color="#1A1A1A")
    ax.set_yticks(ys); ax.set_yticklabels([k for k, _ in items], fontsize=12)
    ax.invert_yaxis()
    ax.set_xlim(0, max(v for _, v in items) * 1.28)
    ax.set_xlabel("Số hồ sơ được xác nhận hợp lệ", fontsize=12)
    ax.grid(axis="x", alpha=0.25, linestyle=":")
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    titles(ax, f"CÁC PHƯƠNG THỨC ĐỐI SOÁT ĐÃ XÁC NHẬN {d.n_ok} HỒ SƠ",
           "Nguồn: cột “Kết quả đối soát MST” — chỉ 34/330 hồ sơ cần thuật toán mờ, "
           "phần còn lại khớp chính xác")
    save(fig, "hinh_5_7_phuong_thuc_doi_soat.png")


def fig_doanh_nghiep_hoc_phan(d):
    top, n_uniq = d.companies(10)
    courses = d.courses()
    fig = plt.figure(figsize=(16.5, 7.6))
    gs = fig.add_gridspec(1, 2, width_ratios=[1.25, 1], wspace=0.45)

    ax1 = fig.add_subplot(gs[0, 0])
    labels = [k[:38] + ("…" if len(k) > 38 else "") for k, _ in top][::-1]
    vals = [v for _, v in top][::-1]
    ax1.barh(range(len(vals)), vals, color=BLUE, height=0.6, edgecolor="#4D4D4D")
    for y, v in zip(range(len(vals)), vals):
        ax1.text(v + 0.8, y, f"{v} ({v/d.n_ok*100:.1f}%)", va="center", fontsize=11.5,
                 fontweight="bold")
    ax1.set_yticks(range(len(vals))); ax1.set_yticklabels(labels, fontsize=11)
    ax1.set_xlim(0, max(vals) * 1.22)
    ax1.set_xlabel("Số giấy giới thiệu đã cấp", fontsize=11.5)
    ax1.grid(axis="x", alpha=0.25, linestyle=":")
    for s in ("top", "right"):
        ax1.spines[s].set_visible(False)
    titles(ax1, "TOP 10 DOANH NGHIỆP TIẾP NHẬN",
           f"330 giấy được cấp cho {n_uniq} doanh nghiệp khác nhau")

    ax2 = fig.add_subplot(gs[0, 1])
    c_items = courses.most_common()
    ax2.bar(range(len(c_items)), [v for _, v in c_items], color="#1E8E5A", width=0.6,
            edgecolor="#4D4D4D")
    for i, (k, v) in enumerate(c_items):
        ax2.text(i, v + 2, str(v), ha="center", fontsize=12, fontweight="bold")
    ax2.set_xticks(range(len(c_items)))
    ax2.set_xticklabels([k for k, _ in c_items], rotation=18, ha="right", fontsize=10.5)
    ax2.set_ylim(0, max(v for _, v in c_items) * 1.2)
    ax2.set_ylabel("Số hồ sơ", fontsize=11.5)
    ax2.grid(axis="y", alpha=0.25, linestyle=":")
    for s in ("top", "right"):
        ax2.spines[s].set_visible(False)
    titles(ax2, "PHÂN BỐ THEO HỌC PHẦN", "Số liệu từ 330 hồ sơ đã cấp giấy")
    save(fig, "hinh_5_8_doanh_nghiep_hoc_phan.png")


def fig_hang_doi():
    fig, ax = plt.subplots(figsize=(16, 8.4))
    ax.set_xlim(0, 1); ax.set_ylim(0, 1); ax.axis("off")
    ax.text(0.5, 0.955, "SƠ ĐỒ HÀNG ĐỢI XỬ LÝ NỀN (BATCH QUEUE) CỦA HỆ THỐNG",
            ha="center", fontsize=17, fontweight="bold", color=NAVY)
    ax.text(0.5, 0.905, "Cho phép chạy hàng trăm hồ sơ mà không vượt giới hạn 6 phút mỗi lần thực thi của Google Apps Script",
            ha="center", fontsize=11.5, color=GREY)

    box(ax, 0.03, 0.66, 0.26, 0.17,
        "1. NGƯỜI DÙNG CHỌN DÒNG\nMenu GGT-CNTT ▸ “In PDF hàng loạt…”\n"
        "(hoặc “Xử lý lại các dòng lỗi / chờ”)", TEAL_L, "#0E7C7B", fs=10.2)
    box(ax, 0.34, 0.66, 0.29, 0.17,
        "2. PHÂN TÍCH DANH SÁCH\nparseRowList_(): 2, 5, 7-10, 15\n"
        "Giới hạn MAX_QUEUE_ROWS = 500 dòng", YELLOW_L, "#C9A227", fs=10.2)
    box(ax, 0.68, 0.66, 0.29, 0.17,
        "3. LƯU HÀNG ĐỢI\nenqueueBatch_(): ghi Script Property\n"
        "GGT_BATCH_QUEUE + trigger mỗi phút", GREEN_L, GREEN, fs=10.2)

    box(ax, 0.68, 0.34, 0.29, 0.20,
        "4. TIẾP TỤC XỬ LÝ\ncontinueBatchQueue_(): khóa script,\n"
        "tối đa MAX_BATCH_ROWS = 25 dòng và\nBATCH_RUNTIME_BUDGET_MS = 240 giây",
        "#DCE9F7", BLUE, fs=10.2)
    box(ax, 0.34, 0.34, 0.29, 0.20,
        "5. CHECKPOINT SAU MỖI DÒNG\nqueue.index / success / failed được ghi lại;\n"
        "nếu execution dừng, lần chạy sau tiếp tục\nđúng dòng còn dang dở",
        RED_L, RED, fs=10.2)
    box(ax, 0.03, 0.34, 0.26, 0.20,
        "6. HOÀN TẤT\nXóa hàng đợi, hủy trigger,\nhiện thông báo (toast) thống kê\n"
        "thành công / lỗi trên Sheet", PURPLE_L, "#6A1B9A", fs=10.2)

    box(ax, 0.03, 0.06, 0.40, 0.16,
        "Giãn cách BATCH_SLEEP_MS = 3.200 ms giữa các hồ sơ\n"
        "để không vượt hạn mức 10 yêu cầu / 30 giây của API Thuế", "#FFFFFF", ORANGE, fs=10.5)
    box(ax, 0.50, 0.06, 0.47, 0.16,
        "cancelBatchQueue(): dừng phần còn lại của lô;\n"
        "mọi tiến độ đã checkpoint vẫn được giữ nguyên", "#FFFFFF", GREY, fs=10.5)

    arrow(ax, (0.29, 0.745), (0.34, 0.745)); arrow(ax, (0.63, 0.745), (0.68, 0.745))
    arrow(ax, (0.825, 0.66), (0.825, 0.545), color=BLUE)
    arrow(ax, (0.68, 0.44), (0.63, 0.44), color=RED)
    arrow(ax, (0.34, 0.44), (0.29, 0.44), color="#6A1B9A")
    arrow(ax, (0.825, 0.34), (0.825, 0.20), color=GREEN, ls="--", rad=0.0)
    ax.text(0.845, 0.265, "còn dòng chưa xử lý → trigger 1 phút chạy tiếp",
            fontsize=10.5, color=GREEN, rotation=0)
    ax.text(0.5, 0.005, "Thực tế đợt 17–18/10/2026: 4 lần chạy nền liên tục, dài nhất 60 phút 51 giây, "
                        "hoàn tất 361 hồ sơ.", ha="center", fontsize=11, color=NAVY)
    save(fig, "hinh_4_5_hang_doi_xu_ly_nen.png")


def fig_trang_thai():
    fig, ax = plt.subplots(figsize=(15.5, 8.0))
    ax.set_xlim(0, 1); ax.set_ylim(0, 1); ax.axis("off")
    ax.text(0.5, 0.955, "MÁY TRẠNG THÁI HỒ SƠ CẤP GIẤY GIỚI THIỆU",
            ha="center", fontsize=17, fontweight="bold", color=NAVY)

    box(ax, 0.03, 0.52, 0.20, 0.20, "MỚI NỘP\n(trạng thái rỗng)", "#E8ECF1", "#5B6B7B", fs=12)
    box(ax, 0.30, 0.52, 0.20, 0.20, "PROCESSING\n(đang xử lý)", YELLOW_L, "#C9A227", fs=12)
    box(ax, 0.63, 0.70, 0.32, 0.17,
        "ĐÃ TẠO PDF / ĐÃ GỬI\n(thành công, có LINK PDF GGT)", GREEN_L, GREEN, fs=12)
    box(ax, 0.63, 0.44, 0.32, 0.17,
        "CHỜ DUYỆT\n(chỉ khi bật REQUIRE_APPROVAL)", "#DCE9F7", BLUE, fs=12)
    box(ax, 0.63, 0.19, 0.32, 0.17,
        "LỖI: <nguyên nhân>\n(sai MST, tên không khớp, SĐT sai…)", RED_L, RED, fs=12)
    box(ax, 0.03, 0.09, 0.26, 0.17, "BỎ QUA KHI CHẠY LẠI\nisAlreadyProcessed_ = true",
        "#E8ECF1", "#5B6B7B", fs=10.5, weight="normal")

    lbl = dict(fontsize=10.5, ha="center", va="center",
               bbox=dict(boxstyle="round,pad=0.28", facecolor="white", edgecolor="none"))

    arrow(ax, (0.235, 0.62), (0.30, 0.62))
    ax.text(0.2675, 0.665, "bắt đầu", color=NAVY, **lbl)
    arrow(ax, (0.505, 0.695), (0.63, 0.775), color=GREEN)
    ax.text(0.551, 0.755, "dữ liệu & MST hợp lệ", color=GREEN, **lbl)
    arrow(ax, (0.505, 0.60), (0.63, 0.535), color=BLUE)
    ax.text(0.551, 0.632, "REQUIRE_APPROVAL = true", color=BLUE, **lbl)
    arrow(ax, (0.505, 0.545), (0.63, 0.33), color=RED)
    ax.text(0.551, 0.468, "sai định dạng /\nkhông khớp", color=RED, **lbl)
    arrow(ax, (0.79, 0.44), (0.79, 0.365), color=ORANGE)
    ax.text(0.80, 0.402, "approveSelectedRow() duyệt", color=ORANGE,
            fontsize=10.5, ha="center", va="center",
            bbox=dict(boxstyle="round,pad=0.28", facecolor="white", edgecolor="none"))
    arrow(ax, (0.632, 0.30), (0.13, 0.515), color=ORANGE, rad=-0.18)
    ax.text(0.36, 0.345, "retryPending() sau khi sinh viên sửa dữ liệu",
            color=ORANGE, fontsize=10.5, ha="center", va="center",
            bbox=dict(boxstyle="round,pad=0.28", facecolor="white", edgecolor="none"))
    ax.text(0.5, 0.02, "Mặc định của bản triển khai: AUTO_NUMBERING = false, REQUIRE_APPROVAL = false, "
                       "SEND_EMAIL = false, AI_ENABLED = false.",
            ha="center", fontsize=11, color=GREY)
    save(fig, "hinh_4_4_may_trang_thai.png")


def fig_kien_truc():
    fig, ax = plt.subplots(figsize=(18, 10.2))
    ax.set_xlim(0, 1); ax.set_ylim(0, 1); ax.axis("off")
    ax.text(0.5, 0.968, "KIẾN TRÚC TỔNG THỂ HỆ THỐNG GGT-CNTT AUTO (Code.gs mới)",
            ha="center", fontsize=17.5, fontweight="bold", color=NAVY)

    box(ax, 0.015, 0.685, 0.225, 0.225,
        "1. TIẾP NHẬN DỮ LIỆU\n• Google Form\n• Sheet “Form Responses 1”\n  (20 cột)\n"
        "• onFormSubmit + LockService", "#DCE9F7", BLUE, fs=10.5)
    box(ax, 0.262, 0.685, 0.225, 0.225,
        "2. CHUẨN HÓA DỮ LIỆU\n• Khử E+ cho\n  MSSV / SĐT / MST\n• Bù số 0 đầu, +84 → 0\n"
        "• Title Case, bỏ dấu", GREEN_L, GREEN, fs=10.5)
    box(ax, 0.509, 0.685, 0.225, 0.225,
        "3. ĐỐI SOÁT THUẾ\n• API xinvoice.vn\n  (cache 6 giờ)\n• Retry 429 + backoff\n"
        "• Weighted Dice, viết tắt", YELLOW_L, "#C9A227", fs=10.5)
    box(ax, 0.756, 0.685, 0.229, 0.225,
        "4. SINH VĂN BẢN\n• Doc mẫu + placeholder\n• trimTrailingBlank_\n"
        "• Xuất PDF 1 trang\n• Xóa Doc tạm", RED_L, RED, fs=10.5)

    box(ax, 0.015, 0.365, 0.30, 0.245,
        "5. GHI KẾT QUẢ VỀ SHEET\n• TRẠNG THÁI GỬI GGT\n  (Đã tạo PDF / Lỗi: …)\n"
        "• LINK PDF GGT\n• THỜI GIAN XỬ LÝ GGT\n• Tên Doanh Nghiệp (theo MST)\n"
        "• Kết quả đối soát MST", PURPLE_L, "#6A1B9A", fs=10.5)
    box(ax, 0.35, 0.365, 0.30, 0.245,
        "6. LƯU TRỮ VÀ PHÂN PHỐI\n• Drive: GGT_Xuat_PDF\n"
        "• PDF_SHARING_MODE =\n  RESTRICTED (mặc định)\n"
        "• SEND_EMAIL = false →\n  mở link ngay trong Sheet", TEAL_L, "#0E7C7B", fs=10.5)
    box(ax, 0.685, 0.365, 0.30, 0.245,
        "7. VẬN HÀNH HÀNG LOẠT\n• Hàng đợi nền:\n  25 dòng/đợt, 240 giây/đợt\n"
        "• Trigger 1 phút chạy tiếp\n• refreshDashboard()\n• retryPending(), menu GGT-CNTT",
        "#FDEBD8", ORANGE, fs=10.5)

    box(ax, 0.015, 0.05, 0.475, 0.22,
        "LỚP MỞ RỘNG (TUỲ CHỌN — ĐANG TẮT)\n"
        "AI_ENABLED = false: chỉ gọi Gemini cho vùng xám\n"
        "0,30 ≤ Dice ≤ 0,50 và độ tin cậy ≥ 0,90;\n"
        "chỉ gửi TÊN DOANH NGHIỆP, không gửi dữ liệu sinh viên.",
        "#F2F2F2", GREY, fs=10.5)
    box(ax, 0.51, 0.05, 0.475, 0.22,
        "KIỂM THỬ TỰ ĐỘNG\n"
        "test_suite.js chạy bằng Node.js: 93/93 ca đạt,\n"
        "phủ chuẩn hóa dữ liệu, đối soát tên, ánh xạ cột,\n"
        "trạng thái, lọc ngày, ghi kết quả và menu onOpen.",
        "#F2F2F2", NAVY, fs=10.5)

    for x in (0.240, 0.487, 0.734):
        arrow(ax, (x, 0.797), (x + 0.021, 0.797))
    arrow(ax, (0.62, 0.683), (0.62, 0.614), color=NAVY)
    arrow(ax, (0.165, 0.681), (0.165, 0.614), color="#6A1B9A")
    arrow(ax, (0.50, 0.614), (0.50, 0.681), color="#0E7C7B", rad=-0.0)
    ax.text(0.5, 0.012, "Toàn bộ giải pháp chạy serverless trên Google Workspace, không cần máy chủ riêng và không phát sinh chi phí hạ tầng.",
            ha="center", fontsize=11.5, color=NAVY)
    save(fig, "hinh_4_1_kien_truc_tong_the.png")


def fig_menu():
    fig, ax = plt.subplots(figsize=(13.5, 9.6))
    ax.set_xlim(0, 1); ax.set_ylim(0, 1); ax.axis("off")
    ax.text(0.5, 0.972, "MENU QUẢN TRỊ “GGT-CNTT” TRÊN GOOGLE SHEET",
            ha="center", fontsize=15.5, fontweight="bold", color=NAVY)

    ax.add_patch(FancyBboxPatch((0.03, 0.885), 0.94, 0.06,
                                boxstyle="round,pad=0.005,rounding_size=0.02",
                                facecolor="#F1F3F4", edgecolor="#9AA0A6"))
    ax.text(0.5, 0.915, "Tệp   Chỉnh sửa   Xem   Chèn   Định dạng   Dữ liệu   Công cụ   Tiện ích mở rộng   [ GGT-CNTT ▾ ]",
            ha="center", fontsize=11.5, color="#3C4043")

    items = ["Kiểm thử dòng cụ thể…",
             "Kiểm thử nhiều dòng…",
             "Kiểm thử dòng cuối cùng",
             "⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯",
             "In PDF hàng loạt…   (chạy hàng đợi nền)",
             "Xử lý lại các dòng lỗi / chờ",
             "Cập nhật Dashboard",
             "⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯⎯",
             "Cài đặt hệ thống (Trigger)"]
    top, h = 0.878, 0.050
    ax.add_patch(Rectangle((0.33, top - len(items) * h - 0.012), 0.50,
                           len(items) * h + 0.028, facecolor="white",
                           edgecolor="#BDC1C6", linewidth=1.4))
    for i, it in enumerate(items):
        y = top - i * h
        if it.startswith("⎯"):
            ax.plot([0.33, 0.83], [y, y], color="#DADCE0", lw=1.2)
            continue
        highlight = i in (1, 4, 6)
        if highlight:
            ax.add_patch(Rectangle((0.33, y - 0.040), 0.50, h - 0.004,
                                   facecolor="#E8F0FE", edgecolor="none"))
        ax.text(0.35, y - 0.015, it, fontsize=10.5, va="center",
                fontweight="bold" if highlight else "normal",
                color=NAVY if highlight else "#202124")

    box(ax, 0.03, 0.02, 0.30, 0.30,
        "CƠ CHẾ CHỐNG QUÁ TẢI\n• MAX_BATCH_ROWS = 25 dòng/đợt\n"
        "• BATCH_RUNTIME_BUDGET_MS\n  = 240.000 (≈ 4 phút)\n"
        "• MAX_QUEUE_ROWS = 500 dòng/lô\n"
        "• BATCH_SLEEP_MS = 3.200 ms",
        "#FDEBD8", ORANGE, fs=9.5)
    box(ax, 0.35, 0.02, 0.30, 0.30,
        "KẾT QUẢ ĐỢT CHẠY THẬT\n• 361/361 hồ sơ được xử lý\n"
        "• 330 tệp PDF thành công\n• 31 hồ sơ bị chặn (có lý do)\n"
        "• 4 đợt nền, dài nhất 60:51",
        GREEN_L, GREEN, fs=9.5)
    box(ax, 0.67, 0.02, 0.30, 0.30,
        "DÀNH CHO CÁN BỘ KHOA\n• Nhập dãy: 2, 5, 7-10, 15\n"
        "• Không cần mở Apps Script\n• Dòng đã cấp giấy: bỏ qua\n"
        "• Dashboard: bấm 1 lần",
        "#DCE9F7", BLUE, fs=9.5)
    save(fig, "hinh_5_5_menu_sheet.png")


CONFIG_CODE = """// ============================================================
// CẤU HÌNH HỆ THỐNG GGT-CNTT AUTO (Nguyễn Hữu Nhân - 522H0045)
// ============================================================
const CONFIG = {
  TEMPLATE_DOC_ID: "1-d1-EzqqMn--gtmMR4dX2fvxXl6Q3zFMb1-…",  // Doc mẫu
  OUTPUT_FOLDER_ID: "1RgRcem14hvJdTa_Cq7As02RnYS8idFQ6",      // Thư mục PDF
  SEND_EMAIL: false,                 // chỉ tạo PDF + ghi link về Sheet
  PDF_SHARING_MODE: "RESTRICTED",    // không mở công khai theo liên kết
  MAX_BATCH_ROWS: 25,                // tối đa 25 dòng mỗi đợt nền
  MAX_QUEUE_ROWS: 500,               // tối đa 500 dòng mỗi lô
  BATCH_RUNTIME_BUDGET_MS: 240000,   // dừng đợt trước mốc 6 phút của GAS
  AI_ENABLED: false,                 // lớp AI tùy chọn, mặc định tắt
  AI_DICE_MIN: 0.30, AI_DICE_MAX: 0.50, AI_CONFIDENCE_THRESHOLD: 0.90,
  AUTO_NUMBERING: false,             // cấp số GGT tự động ({{SO_GGT}})
  REQUIRE_APPROVAL: false,           // bật quy trình “Chờ duyệt”
  FILTER_START_DATE: "2026-08-01",   // mốc lọc đơn theo học kỳ
  TAX_API_URL: "https://api.xinvoice.vn/gdt-api/tax-payer/",
  VERIFY_TAX_CODE: true, USE_OFFICIAL_NAME_ON_DOC: false,
  BATCH_SLEEP_MS: 3200, CACHE_EXPIRY_SEC: 21600
};"""


def fig_config():
    fig, ax = plt.subplots(figsize=(14.5, 8.6))
    ax.set_xlim(0, 1); ax.set_ylim(0, 1); ax.axis("off")
    ax.add_patch(FancyBboxPatch((0.02, 0.02), 0.96, 0.94,
                                boxstyle="round,pad=0.004,rounding_size=0.02",
                                facecolor="#1E1E1E", edgecolor="#3C3C3C"))
    ax.text(0.5, 0.925, "Google Apps Script — Code.gs  (GGT-CNTT Auto, bản cập nhật 10/2026)",
            ha="center", fontsize=12, color="#DDDDDD")
    ax.text(0.05, 0.865, CONFIG_CODE, fontsize=10.4, family="DejaVu Sans Mono",
            color="#D4D4D4", va="top", linespacing=1.65)
    save(fig, "hinh_5_1_apps_script_config.png")


def fig_sheet(d):
    """Minh họa Sheet thật (ẩn danh họ tên sinh viên) với các cột quản trị."""
    def anonym(name):
        parts = str(name).strip().split()
        if len(parts) <= 1:
            return parts[0][:1] + "."
        return " ".join(parts[:-1] + [parts[-1][:1] + "."])

    want = [
        ("khớp chính xác", lambda r: "Khớp chính xác" in d.note(r)),
        ("viết tắt", lambda r: "viết tắt" in d.note(r)),
        ("công ty tổng", lambda r: "công ty tổng" in d.note(r)),
        ("tiếng nước ngoài", lambda r: "tiếng nước ngoài" in d.note(r)),
    ]
    picked, used = [], set()
    for _, cond in want:
        for r in d.success:
            if id(r) in used:
                continue
            if cond(r):
                picked.append(r); used.add(id(r)); break
    for r in d.errors:
        s = d.status_text(r)
        key = "ngừng" if "ngừng hoạt động" in s else ("tồn tại" if "không tồn tại" in s else "khác")
        if key not in used:
            picked.append(r); used.add(key)
        if len(picked) >= 8:
            break
    picked = (picked + [r for r in d.success if id(r) not in used])[:8]

    def short_status(s):
        s = str(s)
        if s.startswith("Lỗi"):
            body = s.split(":", 1)[1].strip() if ":" in s else s
            body = (body.replace("mã số thuế ", "MST ")
                        .replace("thuộc doanh nghiệp", "thuộc DN")
                        .replace("không hợp lệ hoặc ", "không hợp lệ "))
            s = "Lỗi: " + body
        return s[:23] + ("…" if len(s) > 23 else "")

    cols = ["MSSV", "Họ và Tên", "Tên Doanh Nghiệp (SV nhập)", "Mã Số Thuế",
            "TRẠNG THÁI GỬI GGT", "LINK PDF GGT", "THỜI GIAN XỬ LÝ GGT",
            "Kết quả đối soát MST"]
    rows = []
    for r in picked:
        link = str(r[d.idx["LINK PDF GGT"]] or "")
        short = ("…/d/" + link.split("/d/")[1][:11] + "…") if "/d/" in link else "—"
        t = r[d.idx["THỜI GIAN XỬ LÝ GGT"]]
        comp = str(r[d.idx["Tên Doanh Nghiệp"]] or "")
        note = d.note(r)
        rows.append([
            str(r[d.idx["MSSV"]]).replace(".0", ""),
            anonym(r[d.idx["Họ và Tên"]])[:16],
            comp[:26] + ("…" if len(comp) > 26 else ""),
            str(r[d.idx["Mã Số Doanh Nghiệp"]] or "").replace(".0", ""),
            short_status(d.status_text(r)),
            short,
            t.strftime("%d/%m/%Y %H:%M") if isinstance(t, datetime.datetime) else "",
            note[:24] + ("…" if len(note) > 24 else ""),
        ])

    widths = [0.070, 0.085, 0.170, 0.075, 0.150, 0.100, 0.100, 0.150]
    fig, ax = plt.subplots(figsize=(17.5, 6.2))
    ax.axis("off")
    ax.set_title("GOOGLE SHEET SAU KHI HỆ THỐNG XỬ LÝ — TRÍCH 8 DÒNG THỰC TẾ (đã ẩn danh họ tên)",
                 fontsize=14.5, fontweight="bold", color=NAVY, pad=14)
    tbl = ax.table(cellText=rows, colLabels=cols, cellLoc="left", colWidths=widths,
                   bbox=[0.005, 0.0, 0.99, 0.90])
    tbl.auto_set_font_size(False); tbl.set_fontsize(8.2)
    for (row, col), cell in tbl.get_celld().items():
        cell.set_edgecolor("#B7C3CF")
        cell.set_text_props(wrap=False)
        if row == 0:
            cell.set_facecolor("#1F4E79"); cell.set_text_props(color="white", fontweight="bold")
        else:
            txt = rows[row - 1][4]
            cell.set_facecolor("#FFFFFF" if row % 2 else "#F7F9FB")
            if col == 4:
                if txt.startswith("Lỗi"):
                    cell.set_facecolor(RED_L); cell.set_text_props(color=RED, fontweight="bold")
                elif txt == "Đã tạo PDF":
                    cell.set_facecolor(GREEN_L); cell.set_text_props(color=GREEN, fontweight="bold")
    save(fig, "hinh_5_3_google_sheet_result.png")


def main():
    if not os.path.exists(XLSX):
        sys.exit("Không tìm thấy file Sheet: " + XLSX)
    d = Data(XLSX)
    print(f"Dữ liệu: {d.total} hồ sơ | {d.n_ok} thành công | {d.n_err} lỗi | "
          f"{len(d.runs)} đợt chạy | {d.active_seconds/60:.1f} phút chạy thực tế")
    print("Đang sinh hình...")
    fig_ket_qua(d)
    stats = fig_thoi_gian(d)
    print(f"  Thời gian thực tế: {stats['auto_min']:.1f} phút cho {d.total} hồ sơ; "
          f"{stats['per_auto']:.1f} giây/hồ sơ; tiết kiệm {stats['save_pct']:.1f}%")
    fig_dashboard(d)
    fig_doi_soat(d)
    fig_doanh_nghiep_hoc_phan(d)
    fig_hang_doi()
    fig_trang_thai()
    fig_kien_truc()
    fig_menu()
    fig_config()
    fig_sheet(d)
    print("Hoàn tất.")


if __name__ == "__main__":
    main()
